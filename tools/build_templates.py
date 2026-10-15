#!/usr/bin/env python3
"""Собирает рабочие шаблоны скилла из официального шаблона презентации ВШМ СПбГУ.

Исходник — «Шаблон презентации рус.pptx» из папки брендбука (Google Drive,
раздел «Презентация»). Скрипт:
  1. удаляет служебные слайды (пустой, «Резюме», разделители-оглавления шаблона, флаги);
  2. помечает каждый оставшийся слайд ключом каталога (в заметках: «gsom:<key>»);
  3. приводит шрифты к фирменному Inter (в шаблоне местами Euclid Circular / Roboto / Arial);
  4. приводит красный к фирменному #C31632 (в теме шаблона был #BE1E2D);
  5. ужимает крупные растровые изображения (размер файла 50 МБ -> ~15 МБ);
  6. заменяет логотипы шаблона на официальные файлы брендбука (EN-версии подставляются
     при генерации презентации — см. scripts/gsom_deck.py, lang="en").

Запуск:
  python3 tools/build_templates.py "<путь>/Шаблон презентации рус.pptx" "<путь к папке Логотип>"
Результат: skills/gsom-presentation/assets/template.pptx
"""
import hashlib
import io
import os
import re
import sys

from lxml import etree
from PIL import Image
from pptx import Presentation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ASSETS = os.path.join(ROOT, "skills", "gsom-presentation", "assets")
sys.path.insert(0, os.path.join(ROOT, "skills", "gsom-presentation", "scripts"))
from catalog import CATALOG  # noqa: E402  (номер слайда исходника -> ключ)

BRAND_RED = "C31632"
OLD_RED = "BE1E2D"
FONT_ALIASES = {"Euclid Circular A", "Euclid Circular A Light", "Euclid Circular B",
                "Euclid Circular B Light", "Roboto", "Arial", "Calibri", "Calibri Light"}
MAX_SIDE = 1800          # px — ~135 dpi для фото на весь слайд 13.33", для экрана достаточно
MIN_BYTES = 350_000      # трогаем только крупные картинки

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def delete_slide(prs, slide):
    sld_ids = prs.slides._sldIdLst
    for sld_id in list(sld_ids):
        if prs.part.related_part(sld_id.rId) is slide.part:
            prs.part.drop_rel(sld_id.rId)
            sld_ids.remove(sld_id)
            return


def tag_slide(slide, key):
    tf = slide.notes_slide.notes_text_frame
    tf.text = f"gsom:{key}"


def iter_xml_parts(prs):
    seen = set()
    for part in prs.part.package.iter_parts():
        if id(part) in seen:
            continue
        seen.add(id(part))
        if hasattr(part, "_element"):
            yield part


def normalize_fonts_and_colors(prs):
    for part in iter_xml_parts(prs):
        el = part._element
        for tag in ("latin", "cs", "ea"):
            for node in el.iter(f"{{{NS_A}}}{tag}"):
                if node.get("typeface") in FONT_ALIASES:
                    node.set("typeface", "Inter")
        for node in el.iter(f"{{{NS_A}}}srgbClr"):
            if (node.get("val") or "").upper() == OLD_RED:
                node.set("val", BRAND_RED)


def compress_images(prs):
    saved = 0
    seen = set()
    for part in prs.part.package.iter_parts():
        ct = part.content_type
        if id(part) in seen or ct not in ("image/png", "image/jpeg"):
            continue
        seen.add(id(part))
        blob = part.blob
        if len(blob) < MIN_BYTES:
            continue
        im = Image.open(io.BytesIO(blob))
        w, h = im.size
        k = min(1.0, MAX_SIDE / max(w, h))
        if k < 1.0:
            im = im.resize((round(w * k), round(h * k)), Image.LANCZOS)
        out = io.BytesIO()
        if ct == "image/png":
            im.save(out, "PNG", optimize=True)
        else:
            im.convert("RGB").save(out, "JPEG", quality=85, optimize=True)
        if out.tell() < len(blob):
            saved += len(blob) - out.tell()
            part._blob = out.getvalue()
    return saved


def fix_theme_red(prs):
    for part in prs.part.package.iter_parts():
        if part.content_type.endswith("theme+xml"):
            xml = part.blob.decode("utf-8")
            xml = re.sub(OLD_RED, BRAND_RED, xml, flags=re.I)
            # тема хранится как blob-Part, а не XmlPart
            if hasattr(part, "_element"):
                part._element = etree.fromstring(xml.encode("utf-8"))
            else:
                part._blob = xml.encode("utf-8")


def svg_aspect(svg_bytes):
    m = re.search(rb'viewBox="\s*[-\d.]+[\s,]+[-\d.]+[\s,]+([\d.]+)[\s,]+([\d.]+)', svg_bytes)
    return float(m.group(1)) / float(m.group(2))


# Логотипы в исходном шаблоне (md5 svg-файлов) -> официальный файл брендбука.
# В шаблоне логотипы окрашены в #BE1E2D/#51626F; в брендбуке — #C31632/#4A5761.
LOGO_MD5 = {
    "fed0a8ae1f9d05d87138a923d3a95fb9": ("logo-rus.svg", "logo-eng.svg"),           # красный
    "177997d1f6d3549c821545b999ba523f": ("logo-partially-inverted-rus.svg",
                                         "logo-partially-inverted-eng.svg"),        # белый текст
    "8c20dcb8d75256689ab9b5f239f78597": ("logo-grey-rus.svg", "logo-grey-eng.svg"),  # серый
    "fe058d0b69c90eaf81a095bdf4e0bad3": ("logo-inverted-ru.svg", "logo-inverted-en.svg"),  # белый
}
R_NS = {"r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
P_PIC = "{http://schemas.openxmlformats.org/presentationml/2006/main}pic"


def swap_logos(prs, logo_dir, lang):
    """Подменяет логотипы шаблона на файлы брендбука; ширину рамки подгоняет под пропорции."""
    col = 0 if lang == "ru" else 1
    replaced = {}
    for part in prs.part.package.iter_parts():
        if part.content_type != "image/svg+xml":
            continue
        hit = LOGO_MD5.get(hashlib.md5(part.blob).hexdigest())
        if hit:
            new = open(os.path.join(logo_dir, hit[col]), "rb").read()
            part._blob = new
            replaced[part.partname] = svg_aspect(new)
    holders = list(prs.slide_masters)
    holders += [l for m in prs.slide_masters for l in m.slide_layouts] + list(prs.slides)
    fixed = 0
    for holder in holders:
        rels = holder.part.rels
        for pic in holder.shapes._spTree.iter(P_PIC):
            aspect = None
            for rid in etree.ElementBase.xpath(pic, ".//@r:embed", namespaces=R_NS):
                if rid in rels and not rels[rid].is_external:
                    aspect = replaced.get(rels[rid].target_part.partname, aspect)
            if aspect is None:
                continue
            xfrm = pic.find(".//{%s}xfrm" % NS_A)
            off, ext = xfrm.find("{%s}off" % NS_A), xfrm.find("{%s}ext" % NS_A)
            x, cx, cy = int(off.get("x")), int(ext.get("cx")), int(ext.get("cy"))
            new_cx = int(cy * aspect)
            if x + cx / 2 > prs.slide_width / 2:   # логотип справа — держим правый край
                off.set("x", str(x + cx - new_cx))
            ext.set("cx", str(new_cx))
            fixed += 1
    return len(replaced), fixed


def build(src, logo_dir, lang):
    prs = Presentation(src)
    for n, slide in enumerate(list(prs.slides), 1):
        if n in CATALOG:
            tag_slide(slide, CATALOG[n])
        else:
            delete_slide(prs, slide)
    normalize_fonts_and_colors(prs)
    fix_theme_red(prs)
    n_logo, n_pic = swap_logos(prs, logo_dir, lang)
    saved = compress_images(prs)
    out = os.path.join(ASSETS, "template.pptx")
    prs.save(out)
    print(f"{lang.upper()}: {len(prs.slides)} слайдов, логотипов {n_logo} (фигур {n_pic}), "
          f"сжатие −{saved/1e6:.1f} МБ -> {out} ({os.path.getsize(out)/1e6:.1f} МБ)")


def main(src, logo_dir):
    # Один шаблон на оба языка: EN-логотипы подставляет gsom_deck.py при сборке (lang="en").
    build(src, logo_dir, "ru")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
